#!/usr/bin/env python3
"""
Generate Track 1 Methods Description Excel for MVA Hackathon 2026.
Team: Los omikos | Author: JoseVelasco22
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path

def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Track 1 methods"

    header_font = Font(bold=True, size=11)
    required_fill = PatternFill("solid", fgColor="FFC7CE")
    yellow_fill = PatternFill("solid", fgColor="FFFF00")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws["A1"] = "Methods description — MVA Hackathon 2026, Track 1 (Variant Prediction)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Fill out this form once per model/approach submitted."
    ws["A3"] = "Submissions without a methods description may be judged less favorably on Innovation and Scalability."

    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 100

    qa = [
        ("Team name", "Los omikos", True),
        ("Model number (fill out once per model; up to 6)", "1", False),
        ("Please describe your model/approach in detail.",
         "We applied Exomiser v14.0.0 under three orthogonal configurations to verify robustness: "
         "(1) a targeted 22-gene panel of known MVA/microcephaly/SAC genes; "
         "(2) an expanded 51-gene panel adding DNA-repair and checkpoint genes; "
         "(3) a genome-wide run restricted to coding and splice-site variants via variantEffectFilter. "
         "BUB1B ranked first with identical combined score 0.9819 across all three configurations. "
         "We performed forensic validation of four potential limitations: "
         "(a) coarse CNV screening via depth-of-coverage comparison; "
         "(b) rejection of HLA-DRB1 as a false positive driven by MHC hyperpolymorphism; "
         "(c) phase inference by Mendelian functional parsimony; "
         "(d) promoter/UTR scan with no causal regulatory variants.", False),
        ("(required) If commercially-available Generative AI was used, record the provider, the plan or tier, and the relevant setting.",
         "Alibaba Cloud, Qwen 3 Max, commercial tier, standard data-handling policy (no training on customer content). "
         "Used as a methodological reasoning assistant, not as an automated variant caller.", True),
        ("Please state whether the submission file is the automated output of your computational approach (preferred), or if it has undergone downstream manual review and curation.",
         "The submission file is the automated output of Exomiser v14.0.0, with downstream manual curation restricted to evidence verification.", False),
        ("Please describe any downstream manual review in detail.",
         "Manual review consisted of evidence verification: ClinVar confirmation, orthogonal Exomiser validation, coverage and phase inspection in raw VCF.", False),
        ("Please state whether your approach only used publicly available data (preferred) or if proprietary data was also used.",
         "Only publicly available data was used.", False),
        ("Please describe any public data used in detail.",
         "Exomiser 2406_hg38 transcript annotations; 2406_phenotype ontology (HPO, MGI, IMPC, ZFIN); "
         "ClinVar whitelist; gnomAD v2/v4 population frequencies; OMIM:257300 and OMIM:602789; UniProt O60566 (BubR1).", False),
        ("Please describe any proprietary data used in detail.",
         "None.", False),
        ("Is your approach able to output proposed pairs of compound heterozygous candidate variants, or only single candidate variants?",
         "Yes. Exomiser natively outputs compound heterozygous pairs under AUTOSOMAL_RECESSIVE_COMP_HET inheritance mode.", False),
        ("How did you handle secondary or incidental findings, if any?",
         "No secondary findings were reported. HLA-DRB1 and FANCD2 were rejected as artifacts or low-confidence candidates.", False),
        ("Please provide an estimate of run time and cost for your approach.",
         "Total runtime: ~8 minutes on a consumer laptop (WSL2, 6 GB JVM heap, 11 GB RAM). Compute cost: negligible.", False),
        ("Provide a method abstract (up to 500 words). Please include strengths and limitations of your method as well as methodological detail.",
         "We present an orthogonal-validation pipeline for rare-disease variant prioritization, applied to a single WGS proband "
         "with Mosaic Variegated Aneuploidy syndrome 1 (MVA1). The method is built on Exomiser v14.0.0 but goes beyond a single-run "
         "prioritization by executing three independent configurations and requiring convergence across all three. "
         "The primary finding, a BUB1B compound heterozygote (p.Leu737* + p.Asn1002Lys), achieved a combined score of 0.9819 in all three runs.\n\n"
         "Strengths: orthogonal convergence, forensic validation of four failure modes, ClinVar 2-star evidence, "
         "transparent declaration of discordant scores, low computational cost.\n\n"
         "Limitations: phase inferred by parsimony (not read-phased), p.Asn1002Lys remains VUS, "
         "no BAM for rigorous SV/CNV calling, mosaic variants below ~20% VAF may be missed.", False),
    ]

    row = 5
    for question, answer, required in qa:
        ws.cell(row=row, column=1, value=question).font = header_font
        ws.cell(row=row, column=1).alignment = wrap
        if required:
            ws.cell(row=row, column=1).fill = required_fill
        ans_cell = ws.cell(row=row, column=2, value=answer)
        ans_cell.fill = yellow_fill
        ans_cell.alignment = wrap
        row += 1

    out = Path("submission/JoseVelasco22_track1_methods.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Methods Excel generated: {out}")

if __name__ == "__main__":
    main()
